import json
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

files = [
    r"C:\Users\10618071\Downloads\delphinium_sms_pricing_model.xlsx",
    r"C:\Users\10618071\Downloads\1.xlsx",
    r"C:\Users\10618071\Downloads\2.xlsx",
    r"C:\Users\10618071\Downloads\Book3.xlsx",
    r"C:\Users\10618071\Downloads\Book4.xlsx",
    r"C:\Users\10618071\Downloads\Book5.xlsx",
]

out_dir = Path(r"C:\Users\10618071\Projects\prospectus\.tmp_xlsx_parse")
out_dir.mkdir(exist_ok=True)


def get_defined_names(wb):
    names = {}
    try:
        dns = wb.defined_names
        # openpyxl 3.x
        for name in dns:
            defn = dns[name]
            try:
                names[name] = defn.attr_text
            except Exception:
                names[name] = str(defn)
    except Exception as e:
        names = {"_error": str(e)}
    return names


for fp in files:
    p = Path(fp)
    print("\n" + "=" * 80)
    print("FILE:", p.name, "size", p.stat().st_size)
    wb = openpyxl.load_workbook(fp, data_only=False)
    wb_data = openpyxl.load_workbook(fp, data_only=True)
    print("SHEETS:", wb.sheetnames)
    report = {"file": p.name, "sheets": wb.sheetnames, "sheet_details": {}}

    names = get_defined_names(wb)
    report["defined_names"] = names
    print("DEFINED NAMES:", names)

    for sname in wb.sheetnames:
        ws = wb[sname]
        ws_d = wb_data[sname]
        print(
            "\n--- SHEET:",
            sname,
            "dims",
            ws.dimensions,
            "max_row",
            ws.max_row,
            "max_col",
            ws.max_column,
        )
        rows = []
        formulas = []
        max_r = min(ws.max_row or 1, 250)
        max_c = min(ws.max_column or 1, 40)
        for r in range(1, max_r + 1):
            row_cells = []
            empty = True
            for c in range(1, max_c + 1):
                cell = ws.cell(r, c)
                dcell = ws_d.cell(r, c)
                if cell.value is None and dcell.value is None:
                    row_cells.append(None)
                    continue
                empty = False
                entry = {
                    "addr": f"{get_column_letter(c)}{r}",
                    "raw": cell.value,
                    "calc": dcell.value,
                }
                row_cells.append(entry)
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(entry)
            if not empty:
                parts = []
                for e in row_cells:
                    if e is None:
                        continue
                    raw = e["raw"]
                    calc = e["calc"]
                    if isinstance(raw, str) and raw.startswith("="):
                        parts.append(f"{e['addr']}={raw} => {calc}")
                    else:
                        if calc is not None and calc != raw:
                            parts.append(f"{e['addr']}:{raw!r}/{calc!r}")
                        else:
                            parts.append(f"{e['addr']}:{raw!r}")
                line = f"R{r}: " + " | ".join(parts)
                print(line[:2500])
                rows.append(
                    [
                        (
                            {
                                k: (
                                    v
                                    if isinstance(v, (int, float, bool, type(None)))
                                    else str(v)
                                )
                                for k, v in e.items()
                            }
                            if e
                            else None
                        )
                        for e in row_cells
                    ]
                )
        report["sheet_details"][sname] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "formula_count": len(formulas),
            "formulas": [
                {
                    "addr": f["addr"],
                    "raw": f["raw"],
                    "calc": (
                        f["calc"]
                        if isinstance(f["calc"], (int, float, bool, type(None)))
                        else str(f["calc"])
                    ),
                }
                for f in formulas[:800]
            ],
            "rows": rows,
        }
        print(f"  formulas: {len(formulas)}")

    outp = out_dir / (p.stem + "_dump.json")
    with open(outp, "w", encoding="utf-8") as f:
        json.dump(report, f, default=str, indent=2)
    print("Wrote", outp)

print("\nDONE")
